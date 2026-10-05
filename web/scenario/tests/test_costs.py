"""C1: koszty CAPEX/OPEX — czysta kalkulacja (ręczne przykłady), koszty wyniku symulacji, porównanie, xlsx, role."""
from io import BytesIO
from unittest import TestCase as PlainTestCase

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.roles import GROUP_DESIGNER, GROUP_VIEWER
from equipment.models import CostRate, Equipment
from scenario import services
from scenario.compare import compare_columns
from scenario.costs import compute
from scenario.models import Scenario
from twin.models import WarehouseModel

RATES = {"rack_reach": (100, 200), "rack_vna": (300, 400), "rack_shelf": (50, 50), "dock": (1000, 2000),
         "station": (500, 500), "building_m2": (10, 20), "fleet_unit": (9, 9), "fleet_hour": (1, 1),
         "labor_h": (40, 60)}


class ComputeTests(PlainTestCase):
    def _c(self, **over):
        args = dict(rates=RATES, layout={"positions": {"reach": 10, "vna": 0, "shelf": 2}, "docks": 2, "stations": 1,
                                         "area_m2": 100},
                    fleet={"name": "Reach", "units": 3, "purchase": (1000, 1500), "hour": (2, 4)},
                    labor_h_day=16, fleet_busy_h_day=10, work_days=5, volumes={"pallets": 100, "parcels": 0,
                                                                               "orders": 10})
        return compute(**(args | over))

    def test_capex_by_hand(self):
        c = self._c()["capex"]
        # regały 10×100..200 + półki 2×50 + doki 2×1000..2000 + stanowisko 500 + hala 100×10..20 + flota 3×1000..1500
        self.assertEqual((c["low"], c["high"]), (1000 + 100 + 2000 + 500 + 1000 + 3000,
                                                 2000 + 100 + 4000 + 500 + 2000 + 4500))
        self.assertEqual(c["mid"], round((c["low"] + c["high"]) / 2))
        self.assertNotIn("Regały VNA", [i["label"] for i in c["items"]])           # zero miejsc → brak pozycji

    def test_opex_year_and_per_unit(self):
        r = self._c()
        days = 5 * 52
        self.assertEqual(r["days_year"], days)
        # praca 16 h × 260 dni × 40..60 + flota 10 h × 260 × 2..4
        self.assertEqual((r["opex"]["low"], r["opex"]["high"]), (16 * days * 40 + 10 * days * 2,
                                                                 16 * days * 60 + 10 * days * 4))
        self.assertAlmostEqual(r["per_unit"]["pallets"]["low"], r["opex"]["low"] / (100 * days), places=2)
        self.assertNotIn("parcels", r["per_unit"])                                 # zero paczek → brak wskaźnika
        self.assertIn("orders", r["per_unit"])

    def test_fleet_without_catalog_costs_uses_general_rates(self):
        r = self._c(fleet={"name": "wózki", "units": 2, "purchase": None, "hour": None})
        fleet = next(i for i in r["capex"]["items"] if i["label"].startswith("Flota"))
        self.assertEqual((fleet["low"], fleet["high"]), (18, 18))
        self.assertEqual(r["opex"]["items"][1]["low"], 10 * 260 * 1)


class CostViewsTests(TestCase):
    def setUp(self):
        self.designer = User.objects.create_user("proj", password="x")
        self.designer.groups.add(Group.objects.create(name=GROUP_DESIGNER))
        self.viewer = User.objects.create_user("zarzad", password="x")
        self.viewer.groups.add(Group.objects.create(name=GROUP_VIEWER))
        self.client.force_login(self.designer)
        self.client.post(reverse("scenario:create"), {"name": "Rok bazowy"})
        self.sc = Scenario.objects.get()
        self.wm = WarehouseModel.objects.create(name="Hala C1", floor_width_m=80, floor_depth_m=50)
        for i, role in enumerate(("in_container", "in_pallet", "out", "courier")):
            self.wm.features.create(kind="dock", label=f"Dok {i}", dock_role=role, x_m=0, y_m=5 * i, width_m=3.5,
                                    depth_m=4)
        self.wm.racks.create(zone="V", rack_id="001", n_bays=10, n_levels=4, bay_width_cm=270, depth_cm=110,
                             x_m=20, y_m=20)                                   # 120 miejsc reach

    def _run(self, kind="typical"):
        return services.simulate(self.sc.days.get(kind=kind), self.wm, runs=2, user=self.designer)

    def test_run_costs_from_layout_and_default_rates(self):
        c = services.run_costs(self._run())
        items = {i["label"]: i for i in c["capex"]["items"]}
        self.assertEqual(items["Regały standard (reach)"]["qty"], 120)
        self.assertEqual(items["Doki"]["qty"], 4)
        self.assertEqual(items["Hala (budynek)"]["qty"], 4000)
        lo, hi = CostRate.as_dict()["rack_reach"]
        self.assertEqual((items["Regały standard (reach)"]["low"], items["Regały standard (reach)"]["high"]),
                         (round(120 * lo), round(120 * hi)))
        self.assertGreater(c["opex"]["low"], 0)
        self.assertLessEqual(c["capex"]["low"], c["capex"]["high"])

    def test_catalog_fleet_costs_used(self):
        eq = Equipment.objects.get(name="Reach truck 1,6 t / 10 m")
        self.assertEqual(eq.cost_range("cost_purchase", "cost_purchase_max"), (180000.0, 280000.0))
        self.sc.fleet_equipment, self.sc.fleet_units = eq, 3
        self.sc.save()
        c = services.run_costs(self._run())
        fleet = next(i for i in c["capex"]["items"] if i["label"].startswith("Flota"))
        self.assertEqual((fleet["low"], fleet["high"]), (540000, 840000))

    def test_detail_card_xlsx_and_viewer(self):
        run = self._run()
        page = self.client.get(reverse("scenario:detail", args=[self.sc.pk]))
        self.assertContains(page, "Koszty (widełki)")
        self.assertContains(page, reverse("equipment:rates"))
        self.client.force_login(self.viewer)
        page = self.client.get(reverse("scenario:detail", args=[self.sc.pk]))
        self.assertContains(page, "CAPEX razem")
        self.assertNotContains(page, reverse("equipment:rates"))                # Podgląd: wyniki tak, stawki nie
        wb = load_workbook(BytesIO(self.client.get(reverse("scenario:run_xlsx", args=[run.pk])).content))
        rows = [r[0].value for r in wb["Koszty"].iter_rows(min_row=2)]
        self.assertIn("CAPEX razem", rows)
        self.assertTrue(any(str(v).startswith("OPEX na paletę") for v in rows))

    def test_compare_has_cost_rows_and_cheapest(self):
        a, b = self._run(), self._run("peak")
        resp = self.client.get(reverse("scenario:compare") + f"?ids={a.pk}&ids={b.pk}")
        self.assertContains(resp, "CAPEX (środek widełek)")
        self.assertContains(resp, "OPEX na paletę")
        rows = compare_columns([{"costs": {"capex": {"mid": 10}}}, {"costs": {"capex": {"mid": 7}}}])
        capex = next(r for r in rows if r["label"].startswith("CAPEX"))
        self.assertEqual([c["best"] for c in capex["cells"]], [False, True])        # najtańszy wyróżniony

    def test_rates_screen_roles_and_validation(self):
        url = reverse("equipment:rates")
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(self.designer)
        rates = list(CostRate.objects.all())
        post = {"form-TOTAL_FORMS": len(rates), "form-INITIAL_FORMS": len(rates)}
        for i, r in enumerate(rates):
            post |= {f"form-{i}-id": r.pk, f"form-{i}-low": r.low, f"form-{i}-high": r.high}
        post["form-0-low"], post["form-0-high"] = 500, 100                       # od > do
        self.assertContains(self.client.post(url, post), "musi być nie mniejsze")
        post["form-0-low"], post["form-0-high"] = 10, 20
        self.assertEqual(self.client.post(url, post).status_code, 302)
        rates[0].refresh_from_db()
        self.assertEqual((float(rates[0].low), float(rates[0].high)), (10.0, 20.0))

    def test_equipment_form_cost_range_validation(self):
        eq = Equipment.objects.filter(is_system=False).first() or Equipment.objects.create(
            kind="reach", name="Mój reach", speed_loaded_kmh=10, speed_empty_kmh=11, capacity_kg=1500)
        resp = self.client.post(reverse("equipment:edit", args=[eq.pk]), {
            "kind": "reach", "name": "Mój reach", "speed_loaded_kmh": 10, "speed_empty_kmh": 11, "capacity_kg": 1500,
            "lift_curve": "[]", "cost_purchase": 200000, "cost_purchase_max": 100000})
        self.assertContains(resp, "musi być nie mniejsze")

"""S3b w aplikacji: rola doku (pole, migracja, generator, edytor, symulacja), nośność w API layoutu,
pojemność i strefy w wyniku symulacji, tabela porównania, eksport xlsx, role."""
import importlib
import json
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.roles import GROUP_DESIGNER, GROUP_VIEWER
from masterdata.models import ImportLog, Material, StockItem
from scenario import services
from scenario.models import Scenario, ScenarioRun
from twin.design_generator import generate
from twin.layout import guess_dock_role
from twin.models import WarehouseModel


class DockRoleTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("proj", password="x")
        self.user.groups.add(Group.objects.create(name=GROUP_DESIGNER))
        self.client.force_login(self.user)
        self.wm = WarehouseModel.objects.create(name="Hala", floor_width_m=60, floor_depth_m=40)

    def test_migration_heuristic_matches_shared_function(self):
        mig = importlib.import_module("twin.migrations.0004_rola_doku_nosnosc")
        for label in ("Dok kontenerowy 1", "Dok paczek → kontener 2", "Dok paletowy 1", "Dok FTL 3",
                      "Brama busów 1", "Dok wspólny", "Dok IN 4", ""):
            self.assertEqual(mig._guess(label), guess_dock_role(label), label)

    def test_generator_sets_roles(self):
        roles = {f["label"].rsplit(" ", 1)[0]: f.get("dock_role") for f in generate()["features"]
                 if f["kind"] in ("dock", "gate")}
        self.assertEqual(roles["Dok kontenerowy (przenośnik teleskopowy)"], "in_container")
        self.assertEqual(roles["Dok FTL"], "out")
        self.assertEqual(roles["Dok paczek → kontener (przenośnik teleskopowy)"], "courier")

    def test_layout_api_roundtrip_and_keep_when_absent(self):
        f = self.wm.features.create(kind="dock", label="Dok FTL 1", dock_role="in_container", x_m=0, y_m=0,
                                    width_m=3.5, depth_m=2)
        r = self.wm.racks.create(zone="V", rack_id="001", n_bays=4, n_levels=3, x_m=10, y_m=10, load_kg=1200)
        data = self.client.get(reverse("twin:warehouse_layout_json", args=[self.wm.pk])).json()
        self.assertEqual(data["features"][0]["dock_role"], "in_container")
        self.assertEqual(data["racks"][0]["load_kg"], 1200)
        data["features"][0].pop("dock_role")                      # stary klient bez pola — nie kasuje roli
        data["racks"][0]["load_kg"] = 800
        resp = self.client.post(reverse("twin:warehouse_layout_save", args=[self.wm.pk]), json.dumps(data),
                                content_type="application/json")
        self.assertEqual(resp.status_code, 200, resp.content)
        f.refresh_from_db()
        r.refresh_from_db()
        self.assertEqual((f.dock_role, r.load_kg), ("in_container", 800))
        data = resp.json()
        data["features"][0]["dock_role"] = "nieznana"
        bad = self.client.post(reverse("twin:warehouse_layout_save", args=[self.wm.pk]), json.dumps(data),
                               content_type="application/json")
        self.assertEqual(bad.status_code, 400)

    def test_editor_config_has_dock_roles(self):
        page = self.client.get(reverse("twin:warehouse_layout_editor", args=[self.wm.pk]))
        self.assertContains(page, '"dockRoles": {"in_container"')

    def test_features_form_saves_role_only_for_docks(self):
        url = reverse("twin:warehouse_model_features", args=[self.wm.pk])
        post = {"row_id": ["", ""], "kind": ["dock", "staging"], "label": ["Dok A", "Bufor"],
                "dock_role": ["courier", "courier"], "zone_code": ["", ""], "x_m": ["0", "5"], "y_m": ["0", "5"],
                "width_m": ["3", "8"], "depth_m": ["2", "6"], "angle_deg": ["0", "0"],
                "color_hex": ["", ""], "notes": ["", ""], "deleted_ids": ""}
        self.client.post(url, post)
        self.assertEqual(dict(self.wm.features.values_list("kind", "dock_role")), {"dock": "courier", "staging": ""})


class SimS3bTests(TestCase):
    def setUp(self):
        self.designer = User.objects.create_user("proj", password="x")
        self.designer.groups.add(Group.objects.create(name=GROUP_DESIGNER))
        self.viewer = User.objects.create_user("zarzad", password="x")
        self.viewer.groups.add(Group.objects.create(name=GROUP_VIEWER))
        self.client.force_login(self.designer)
        self.client.post(reverse("scenario:create"), {"name": "Rok bazowy"})
        self.sc = Scenario.objects.get()
        self.wm = WarehouseModel.objects.create(name="Hala S3b", floor_width_m=80, floor_depth_m=50)
        for i, role in enumerate(("in_container", "in_pallet", "out", "out", "courier")):
            self.wm.features.create(kind="dock", label=f"Dok {i}", dock_role=role, x_m=0, y_m=5 * i, width_m=3.5,
                                    depth_m=4)
        self.wm.features.create(kind="zone_adr", label="ADR", x_m=19, y_m=19, width_m=30, depth_m=4)
        self.wm.racks.create(zone="V", rack_id="001", n_bays=10, n_levels=4, bay_width_cm=270, depth_cm=110,
                             x_m=20, y_m=20)                                   # 120 miejsc, w strefie ADR
        self.wm.racks.create(zone="V", rack_id="002", n_bays=10, n_levels=4, bay_width_cm=270, depth_cm=110,
                             x_m=20, y_m=30)
        log = ImportLog.objects.create(kind="stock", name="stan")
        Material.objects.create(code="M-ADR", cartons_per_layer=6, layers_per_pallet=5, adr=True)    # 30 kart.
        Material.objects.create(code="M-ZWYKLY", cartons_per_pallet=50)
        StockItem.objects.bulk_create([StockItem(log=log, location_code=f"L{i}", material_code="M-ADR")
                                       for i in range(150)]
                                      + [StockItem(log=log, location_code=f"K{i}", material_code="M-ZWYKLY")
                                         for i in range(50)])

    def _run(self):
        day = self.sc.days.get(kind="typical")
        return services.simulate(day, self.wm, runs=2, user=self.designer)

    def test_result_has_placement_cpp_and_explicit_roles(self):
        run = self._run()
        cap = run.result["placement"]["capacity"]
        self.assertEqual((cap["positions"], cap["need"]), (240, 200))
        adr = next(z for z in cap["zones"] if z["kind"] == "zone_adr")
        self.assertEqual((adr["positions"], adr["need"]), (120, 150))
        self.assertTrue(any("ADR" in b["problem"] for b in run.result["bottlenecks"]))
        self.assertEqual(run.result["places"]["warnings"], [])              # role jawne — bez zgadywania
        self.assertEqual(run.result["cpp"]["value"], 35.0)                  # (150×30 + 50×50) / 200
        self.assertIn("master data", run.result["cpp"]["source"])
        page = self.client.get(reverse("scenario:detail", args=[self.sc.pk]))
        self.assertContains(page, "Pojemność i rozmieszczenie")
        self.assertContains(page, reverse("scenario:run_xlsx", args=[run.pk]))
        self.assertContains(page, 'data-fs-target="sim-box-')

    def test_xlsx_sheets_and_no_material_lists(self):
        run = self._run()
        self.client.force_login(self.viewer)                                 # Podgląd: wyniki — tak
        resp = self.client.get(reverse("scenario:run_xlsx", args=[run.pk]))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Założenia", "KPI", "Wąskie gardła", "Obsada", "Pojemność", "Koszty"])
        self.assertEqual(wb["Założenia"]["B2"].value, "Rok bazowy")          # wiersz 1 = nagłówek
        self.assertEqual(wb["KPI"]["A1"].value, "Wskaźnik")
        self.assertEqual(wb["Pojemność"]["B2"].value, 240)
        text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
        self.assertNotIn("M-ADR", text)                                      # #25: bez danych źródłowych

    def test_compare_view_and_xlsx(self):
        a = self._run()
        self.wm.racks.create(zone="V", rack_id="003", n_bays=10, n_levels=4, bay_width_cm=270, depth_cm=110,
                             x_m=20, y_m=40)
        b = self._run()
        self.client.force_login(self.viewer)
        url = reverse("scenario:compare") + f"?ids={a.pk}&ids={b.pk}"
        page = self.client.get(url)
        self.assertEqual(page.status_code, 200)
        rows = {r["label"]: r["cells"] for r in page.context["rows"]}
        cap = rows["Miejsca paletowe w layoucie"]
        self.assertEqual([c["value"] for c in cap], [240, 360])
        self.assertTrue(cap[1]["best"])
        self.assertEqual(cap[1]["delta"], 50)
        self.assertContains(page, 'id="cmp-fs"')
        wb = load_workbook(BytesIO(self.client.get(url + "&xlsx=1").content))
        self.assertEqual(wb.sheetnames, ["Porównanie"])
        self.assertEqual(wb["Porównanie"]["B2"].value, 240)

    def test_compare_empty_and_limit(self):
        self.assertContains(self.client.get(reverse("scenario:compare")), "Nie ma jeszcze wyników")
        run = self._run()
        ids = "&".join(f"ids={run.pk}" for _ in range(3))
        page = self.client.get(reverse("scenario:compare") + "?" + ids)
        self.assertEqual(len(page.context["chosen"]), 1)                     # duplikaty zwijane
        self.assertEqual(ScenarioRun.objects.count(), 1)

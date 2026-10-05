"""Eksport wyników symulacji do xlsx (E8): założenia, KPI, wąskie gardła, obsada, pojemność — i tabela porównania.

Zgodnie z decyzją #25 w pliku są tylko założenia scenariusza i wyniki — bez list materiałów i stanów.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .inbound import _hhmm
from .sim.report import KPI_SPEC
from .staffing import PROCESSES, process_hours, staffing

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
NUM = "#,##0.0"


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(header, 1):
        width = (widths or {}).get(i) or min(60, max(10, len(str(h)) + 2,
                                                   *(len(str(r[i - 1])) + 2 for r in rows if len(r) >= i)))
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = NUM
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 50)
    ws.freeze_panes = "A2"
    return ws


def _bytes(wb):
    del wb["Sheet"]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_workbook(run, day):
    """run: ScenarioRun; day: ScenarioDay tego wyniku (do obsady z założeń)."""
    sc, r = run.scenario, run.result
    wb = Workbook()
    cpp = r.get("cpp") or {}
    info = [("Scenariusz", sc.name), ("Dzień", run.get_day_kind_display()), ("Model hali", str(run.model)),
            ("Mnożnik wzrostu", sc.growth), ("Ziarno", run.seed), ("Przebiegów", run.runs),
            ("Kartonów na paletę (kontener)", cpp.get("value")), ("Źródło kartonów/paletę", cpp.get("source", "norma")),
            ("Data symulacji", run.created_at.strftime("%Y-%m-%d %H:%M"))]
    info += [(sc._meta.get_field(f).verbose_name, getattr(sc, f)) for f in [*sc.NORM_FIELDS, *sc.FLEET_FIELDS]]
    _sheet(wb, "Założenia", ["Parametr", "Wartość"], info, {1: 48, 2: 30})

    agg = r.get("agg", {})
    _sheet(wb, "KPI", ["Wskaźnik", "Jednostka", "Średnio", "Najgorszy (P95)"],
           [(a["label"], a["unit"], a["mean"], a["worst"]) for k, *_ in KPI_SPEC if (a := agg.get(k))], {1: 40})

    _sheet(wb, "Wąskie gardła", ["Waga", "Obszar", "Miejsce", "Okno", "Problem", "Podpowiedź"],
           [("krytyczne" if b["severity"] == "error" else "ostrzeżenie", b.get("area", ""), b.get("where", ""),
             b.get("window", ""), b.get("problem", ""), b.get("suggestion", "")) for b in r.get("bottlenecks", [])],
           {5: 70, 6: 60})

    shifts = sc.shift_dicts()
    plans = {lvl: staffing(process_hours(day.demand(lvl), day.outbound_demand(lvl)), shifts) for lvl in ("avg", "max")}
    rows = []
    for i, (proc, label) in enumerate(PROCESSES):
        util = (agg.get(f"util_{proc}") or {}).get("mean")
        for j, sh in enumerate(plans["avg"][i]["shifts"]):
            rows.append((label, f"{_hhmm(sh['start_h'])}–{_hhmm(sh['end_h'])}", sh["break_min"], sh["people"],
                         sh["needed"], plans["max"][i]["shifts"][j]["needed"], util))
    _sheet(wb, "Obsada", ["Proces", "Zmiana", "Przerwa [min]", "Zakładana", "Potrzebna (średnio)",
                          "Potrzebna (max)", "Wykorzystanie w symulacji [%]"], rows, {1: 30})

    cap = (r.get("placement") or {}).get("capacity") or {}
    prow = [("Miejsca paletowe w layoucie", cap.get("positions")), ("Potrzeba (stan × wzrost)", cap.get("need")),
            ("Wypełnienie [%]", cap.get("fill_pct")), ("Ciężkie palety (najwyższa klasa wagi)", cap.get("heavy_need")),
            ("Miejsca na dolnych poziomach", cap.get("heavy_low_positions"))]
    prow += [(f"Strefa: {z['label']} — miejsca / potrzeba", f"{z['positions']} / {z['need']}")
             for z in cap.get("zones", [])]
    _sheet(wb, "Pojemność", ["Pozycja", "Wartość"], prow, {1: 50, 2: 20})
    return _bytes(wb)


def compare_workbook(columns, rows):
    """columns: [nagłówek kolumny]; rows: `compare.compare_columns`."""
    wb = Workbook()
    data = []
    for row in rows:
        data.append([row["label"] + (f" [{row['unit']}]" if row["unit"] else ""),
                     *[c["value"] for c in row["cells"]]])
    ws = _sheet(wb, "Porównanie", ["Wskaźnik", *columns], data, {1: 40})
    for i, row in enumerate(rows, 2):
        for j, c in enumerate(row["cells"], 2):
            if c["best"]:
                ws.cell(i, j).font = Font(bold=True, color="166534")
    return _bytes(wb)

"""K2: osprzęt (wpływ na udźwig i czas), import własnych modeli z xlsx, nowe klasy ogólne, role."""
import io
from unittest import TestCase as PlainTestCase

import openpyxl
from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from core.roles import GROUP_DESIGNER, GROUP_VIEWER
from equipment import importers
from equipment.catalog import apply_attachments, capacity_at
from equipment.models import Equipment


class AttachmentTests(PlainTestCase):
    P = {"capacity_kg": 2000, "lift_curve": [[6, 1600], [10, 1000]], "pick_s": 20, "drop_s": 25}

    def test_reduction_and_time(self):
        out = apply_attachments(self.P, ["sideshift", "rotator"])          # −5 % −15 %, −3 s +8 s
        self.assertEqual(out["capacity_kg"], 1600)
        self.assertEqual(out["lift_curve"], [[6, 1280], [10, 800]])
        self.assertEqual((out["pick_s"], out["drop_s"]), (25, 30))
        self.assertEqual(capacity_at(out["capacity_kg"], out["lift_curve"], 8), 1040)
        self.assertEqual(self.P["capacity_kg"], 2000)                      # wejście nietknięte

    def test_none_or_unknown_is_identity(self):
        self.assertEqual(apply_attachments(self.P, []), self.P)
        self.assertEqual(apply_attachments(self.P, ["nieznany"]), self.P)


class ParseTests(PlainTestCase):
    COLS = importers._cols(["Typ", "Nazwa", "Producent", "Udźwig [kg]", "Maks. podnoszenie [m]",
                            "Krzywa udźwigu (wys:kg;…)", "Osprzęt"])

    def test_row_ok(self):
        r = importers.parse_row(["Reach truck (wysokiego składowania)", "Model testowy A", "Firma X", "1 600",
                                 "10,5", "6:1600;10:1000", "Przesuw boczny; Obrotnica"], self.COLS)
        self.assertEqual((r["kind"], r["capacity_kg"], r["max_lift_m"]), ("reach", 1600, 10.5))
        self.assertEqual(r["lift_curve"], [[6.0, 1600], [10.0, 1000]])
        self.assertEqual(r["attachments"], ["sideshift", "rotator"])
        self.assertEqual(r["manufacturer"], "Firma X")

    def test_row_errors(self):
        for row, msg in [(["Rakieta", "M", "", 1000], "nieznany typ"), (["vna", "", "", 1000], "brak nazwy"),
                         (["vna", "M", "", ""], "brak udźwigu"), (["vna", "M", "", "abc"], "nie liczba"),
                         (["vna", "M", "", 900, "", "", "Hak"], "nieznany osprzęt")]:
            with self.subTest(msg=msg), self.assertRaisesRegex(ValueError, msg):
                importers.parse_row(row, self.COLS)


def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("proj", password="x")
        self.user.groups.add(Group.objects.create(name=GROUP_DESIGNER))
        self.viewer = User.objects.create_user("zarzad", password="x")
        self.viewer.groups.add(Group.objects.create(name=GROUP_VIEWER))
        self.client.force_login(self.user)

    def _post(self, rows):
        f = SimpleUploadedFile("katalog.xlsx", _xlsx(rows))
        return self.client.post(reverse("equipment:import"), {"plik": f}, follow=True)

    def test_new_market_classes_exist(self):
        for kind in ("stacker", "order_picker", "tractor"):
            self.assertTrue(Equipment.objects.filter(is_system=True, kind=kind).exists())
        self.assertTrue(Equipment.objects.filter(is_system=True, kind="vna", max_lift_m=18).exists())

    def test_import_creates_updates_and_fills_from_class(self):
        head = ["Typ", "Nazwa", "Producent", "Udźwig [kg]", "Maks. podnoszenie [m]", "Osprzęt"]
        resp = self._post([head, ["reach", "Model testowy A", "Firma X", 1600, 10, "Przesuw boczny"],
                           ["Rakieta", "Zły", "", 1, 1, ""]])
        self.assertContains(resp, "1 nowych, 0 zaktualizowanych")
        self.assertContains(resp, "Wiersz 3: nieznany typ")
        m = Equipment.objects.get(name="Model testowy A")
        self.assertEqual((m.is_system, m.manufacturer, m.attachments), (False, "Firma X", ["sideshift"]))
        self.assertIsNotNone(m.aisle_m)                                   # z klasy ogólnej reach
        self.assertIn("Z klasy ogólnej", m.notes)
        self.assertEqual(m.params()["capacity_kg"], 1520)                 # przesuw boczny −5 %
        resp = self._post([head, ["reach", "Model testowy A", "Firma X", 1800, 10, ""]])
        self.assertContains(resp, "0 nowych, 1 zaktualizowanych")
        m.refresh_from_db()
        self.assertEqual((m.capacity_kg, m.attachments), (1800, []))
        page = self.client.get(reverse("equipment:list") + "?producent=Firma X")
        self.assertContains(page, "Model testowy A")
        self.assertNotContains(page, "Reach truck 1,6 t / 10 m")

    def test_system_class_untouched(self):
        sys_name = "Reach truck 1,6 t / 10 m"
        resp = self._post([["Typ", "Nazwa", "Udźwig [kg]"], ["reach", sys_name, 9999]])
        self.assertContains(resp, "nazwa klasy systemowej")
        self.assertEqual(Equipment.objects.get(name=sys_name).capacity_kg, 1600)

    def test_template_and_roles(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.client.get(reverse("equipment:import_template")).content))
        self.assertEqual(wb.active.cell(1, 1).value, "Typ")
        cols = importers._cols([c.value for c in wb.active[1]])
        self.assertEqual(importers.parse_row([c.value for c in wb.active[2]], cols)["name"], "Model testowy A")
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("equipment:import_template")).status_code, 403)
        f = SimpleUploadedFile("k.xlsx", _xlsx([["Typ", "Nazwa"]]))
        self.assertEqual(self.client.post(reverse("equipment:import"), {"plik": f}).status_code, 403)
